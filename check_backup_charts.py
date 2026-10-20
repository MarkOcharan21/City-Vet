from pptx import Presentation

prs = Presentation('frontend/SNED-Capstone-Defense.backup.pptx')

for slide_idx in [8, 9, 10]:
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.has_chart:
            chart = shape.chart
            cats = []
            if chart.plots and chart.plots[0].categories:
                cats = [str(c) for c in chart.plots[0].categories]
            series_list = []
            if chart.plots:
                for series in chart.plots[0].series:
                    vals = [v for v in series.values] if series.values else []
                    series_list.append((series.name, vals))
            title = chart.chart_title.text_frame.text if chart.has_title else None
            has_legend = chart.has_legend
            legend_pos = chart.legend.position if has_legend else None
            
            print(f'Slide {slide_idx+1}:')
            print(f'  Categories: {cats}')
            for name, vals in series_list:
                print(f'  Series: {name} = {vals}')
            print(f'  Title: {title}')
            print(f'  Legend: {has_legend}, pos={legend_pos}')