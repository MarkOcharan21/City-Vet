from pptx import Presentation

prs = Presentation('City Vet.pptx')

def extract_chart_info(chart):
    info = {
        'chart_type': chart.chart_type,
        'has_legend': chart.has_legend,
        'legend_position': None,
        'legend_include_in_layout': None,
        'series': [],
        'categories': [],
        'title': None,
    }
    
    if chart.has_legend:
        info['legend_position'] = chart.legend.position
        info['legend_include_in_layout'] = chart.legend.include_in_layout
    
    if chart.has_title:
        info['title'] = chart.chart_title.text_frame.text
    
    if chart.plots:
        plot = chart.plots[0]
        if plot.categories:
            info['categories'] = [str(c) for c in plot.categories]
        for series in plot.series:
            s_info = {
                'name': series.name,
                'values': [v for v in series.values] if series.values else [],
            }
            info['series'].append(s_info)
    
    return info

for slide_idx in [8, 9, 10]:
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.has_chart:
            info = extract_chart_info(shape.chart)
            print(f"Slide {slide_idx+1}:")
            print(f"  Type: {info['chart_type']}")
            print(f"  Categories: {info['categories']}")
            for s in info['series']:
                print(f"  Series: {s['name']} = {s['values']}")
            print(f"  Legend: {info['has_legend']}, pos={info['legend_position']}")
            print(f"  Title: {info['title']}")