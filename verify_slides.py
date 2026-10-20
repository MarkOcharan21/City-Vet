from pptx import Presentation
prs = Presentation('City Vet.pptx')
for i in [8, 9, 10]:
    slide = prs.slides[i]
    print(f'=== Slide {i+1} ===')
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                print(f'  Shape {j}: "{text[:100]}"')
        elif shape.has_chart:
            print(f'  Shape {j}: Chart - {shape.chart.chart_type}')