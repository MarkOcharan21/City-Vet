from pptx import Presentation
import os

for fname in ['City Vet.pptx', 'SNED-Capstone-Defense (1).pptx']:
    if not os.path.exists(fname):
        print(f'NOT FOUND: {fname}')
        continue
    prs = Presentation(fname)
    print(f'=== {fname} ===')
    print(f'Slides: {len(prs.slides)}')
    for i, slide in enumerate(prs.slides):
        print(f'  Slide {i+1}: {len(slide.shapes)} shapes')
        for j, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                text = shape.text_frame.text[:80]
                print(f'    Shape {j}: {shape.shape_type} - "{text}"')
            elif shape.shape_type == 13:
                print(f'    Shape {j}: Picture {shape.width}, {shape.height}')
            elif shape.has_chart:
                print(f'    Shape {j}: Chart - {shape.chart.chart_type}')
            else:
                print(f'    Shape {j}: {shape.shape_type}')